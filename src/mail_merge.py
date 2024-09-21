import os
import pandas as pd
from docx import Document
from .config import EMAIL_LIST_PATH, MAIL_MERGE_TEMPLATE_PATH, MERGED_FOLDER

def mail_merge(template_path=MAIL_MERGE_TEMPLATE_PATH, excel_data_path=EMAIL_LIST_PATH, output_folder=MERGED_FOLDER):
    """
    Performs mail merge by populating a Word template with data from an Excel file.
    """
    # Ensure output folder exists
    os.makedirs(output_folder, exist_ok=True)

    # Load Excel data
    df = pd.read_excel(excel_data_path, sheet_name='Email_List')

    for _, row in df.iterrows():
        recipient_name = row['Recipient Name']
        document = Document(template_path)

        # Replace placeholders in the document
        for paragraph in document.paragraphs:
            if '[[Recipient]]' in paragraph.text:
                paragraph.text = paragraph.text.replace('[[Recipient]]', recipient_name)

        # Save the personalized document
        output_path = os.path.join(output_folder, f"{recipient_name}_letter.docx")
        document.save(output_path)
        print(f"Mail merge document created for {recipient_name} at {output_path}")
